"""Bulk import for the rental inventory — spreadsheet in, rental items out.

The importer owns no schema of its own. Every row is turned into the same
dict the Add Rental Item form posts and handed to `content._clean_rental`,
so a rule that holds for one item typed by hand holds for five hundred
arriving in a spreadsheet. Pictures go through `media.ingest_library_image`,
the one function the panel's own uploader uses, so they are sniffed, opened,
re-encoded to WebP and content-addressed exactly as a hand upload is.

The flow is deliberately two-phase. An upload only *stages* a job: the file
is parsed, every row validated and every picture resolved, and the admin is
shown what would happen. Nothing is written until a second call names that
staged job. Staged jobs live in memory, expire, and take their temporary
copy of the upload with them.
"""
from __future__ import annotations

import csv
import io
import ipaddress
import json
import re
import secrets
import socket
import threading
import time
import urllib.parse
import zipfile

from . import config

# ---------------- limits ----------------

MAX_ROWS = 2000              # a spreadsheet longer than this is a mistake
MAX_SHEET_BYTES = 12 * 1024 * 1024
MAX_ZIP_BYTES = 96 * 1024 * 1024
MAX_ZIP_ENTRIES = 4000
MAX_ZIP_UNCOMPRESSED = 320 * 1024 * 1024   # zip-bomb ceiling
MAX_ZIP_RATIO = 200                        # per-entry compression ratio ceiling
IMAGE_URL_TIMEOUT_S = 20
MAX_IMAGE_BYTES = 15 * 1024 * 1024         # matches media.LIBRARY_MAX_BYTES
IMAGE_COLUMNS = 5                          # image_1 … image_5
STAGE_TTL_S = 30 * 60
BATCH_SIZE = 25
KEEP_HISTORY = 50

CLEAR_TOKEN = "[clear]"      # the one way a spreadsheet may empty a field


class ImportError_(Exception):
    """User-facing message; safe to show in the panel."""


# ---------------- the columns, derived from the rental item itself ----------------

#: (column, label, help) — the importable half of a rental item. The names are
#: the business names an admin sees in the form, not internal keys.
COLUMNS: list[tuple[str, str, str]] = [
    ("id", "Required", "Permanent item id: lowercase letters, digits and dashes, "
                       "e.g. rent-led-wall. This is what identifies the item on re-import."),
    ("code", "Optional", "Your own reference or SKU, e.g. EM-R-001. Up to 40 characters."),
    ("name", "Required", "Item name shown on the rental card. Up to 160 characters."),
    ("category", "Required", "Must match an existing category exactly, unless you tick "
                             "“Create missing categories” when importing."),
    ("description", "Optional", "Up to 2000 characters. Plain text — HTML is removed."),
    ("specifications", "Optional", "One specification per line, or separated by | . "
                                   "Up to 15, each up to 200 characters."),
    ("search_tags", "Optional", "Comma separated. Up to 12, each up to 40 characters."),
    ("stock_ksa", "Optional", "Whole number 0–100000. Blank means 0 on a new item."),
    ("stock_uae", "Optional", "Whole number 0–100000. Blank means 0 on a new item."),
    ("featured", "Optional", "yes or no. Featured items are shown first."),
] + [
    (f"image_{i}", "Optional",
     ("Primary image — the picture on the rental card. " if i == 1 else f"Image {i}. ") +
     "A https:// URL, a file name inside the images ZIP, or an existing "
     "/media/… or /assets/… path.")
    for i in range(1, IMAGE_COLUMNS + 1)
]

COLUMN_NAMES = [c[0] for c in COLUMNS]

#: spellings an admin might reasonably use for the same column
_ALIASES = {
    "item_id": "id", "itemid": "id", "slug": "id", "rental_id": "id",
    "sku": "code", "item_code": "code", "reference": "code",
    "title": "name", "item_name": "name", "product_name": "name",
    "category_name": "category", "type": "category",
    "desc": "description", "details": "description",
    "specs": "specifications", "specification": "specifications",
    "tags": "search_tags", "keywords": "search_tags", "search tags": "search_tags",
    "quantity_ksa": "stock_ksa", "qty_ksa": "stock_ksa", "ksa": "stock_ksa",
    "stock_saudi": "stock_ksa", "saudi_stock": "stock_ksa",
    "quantity_uae": "stock_uae", "qty_uae": "stock_uae", "uae": "stock_uae",
    "is_featured": "featured", "feature": "featured",
    "primary_image": "image_1", "main_image": "image_1", "image": "image_1",
}


def _norm_header(raw: str) -> str:
    key = re.sub(r"[\s\-]+", "_", str(raw or "").strip().lower())
    key = re.sub(r"[^a-z0-9_]", "", key)
    if key in _ALIASES:
        return _ALIASES[key]
    return key


_TRUE = {"yes", "y", "true", "1", "on", "featured"}
_FALSE = {"no", "n", "false", "0", "off", ""}


# ---------------- spreadsheet writing (template, error report) ----------------

_FORMULA_LEAD = ("=", "+", "-", "@", "\t", "\r")


def csv_safe(value) -> str:
    """A cell that Excel will not execute. A value that begins like a formula
    is prefixed with an apostrophe — the standard defence, and visible, so
    nobody wonders later why the sheet looks odd."""
    text = "" if value is None else str(value)
    if text[:1] in _FORMULA_LEAD:
        return "'" + text
    return text


def _example_row() -> dict:
    return {
        "id": "rent-led-wall-3x2", "code": "EM-R-101",
        "name": "3x2m LED Video Wall", "category": "Displays & LED",
        "description": "Indoor LED wall on a ground-supported frame, delivered, "
                       "installed and tested before the doors open.",
        "specifications": "3m x 2m active area|3.9mm pixel pitch|Ground support frame included",
        "search_tags": "led, video wall, screen", "stock_ksa": "4", "stock_uae": "2",
        "featured": "yes",
        "image_1": "https://example.com/led-wall-front.jpg",
        "image_2": "led-wall-side.jpg", "image_3": "", "image_4": "", "image_5": "",
    }


def _instructions_rows(categories: list[str]) -> list[list[str]]:
    rows = [
        ["Elite Marcom — Rental Items bulk import"],
        [],
        ["Fill in the Items sheet. One rental item per row. Keep the header row as it is."],
        ["The importer shows you a preview before anything is created, so a mistake here "
         "costs nothing."],
        [],
        ["Column", "Required?", "What to put in it"],
    ]
    rows += [[name, need, help_] for name, need, help_ in COLUMNS]
    rows += [
        [],
        ["Categories"],
        ["A category must already exist, or the row is reported as an error. Tick "
         "“Create missing categories” on the import screen to allow new ones."],
        ["Categories that exist today:"],
    ]
    rows += [[c] for c in categories] or [["(none yet)"]]
    rows += [
        [],
        ["Images"],
        ["Three ways to give a picture, in any image_ column:"],
        ["1. A https:// address. The importer downloads it and stores it in the Media "
         "Library. The site never hotlinks someone else's server."],
        ["2. A file name, if you also upload an images ZIP — for example CHR-001-1.jpg."],
        ["3. An existing path already on the site, such as /media/abc123.webp."],
        ["JPG, PNG and WebP only. Up to 15 MB each. Every picture is converted to WebP."],
        ["image_1 is the primary picture — the one on the rental card."],
        [],
        ["Duplicates"],
        ["Items are matched on the id column. If an item with that id already exists you "
         "choose, on the import screen, whether to skip it or update it."],
        ["When updating, a blank cell leaves the existing value alone. To empty a field "
         f"on purpose, put {CLEAR_TOKEN} in the cell."],
        [],
        ["Numbers"],
        ["stock_ksa and stock_uae are whole numbers from 0 to 100000."],
        [],
        ["Example"],
        ["The Items sheet already contains one example row. Delete it before importing, "
         "or leave it and skip that row in the preview."],
    ]
    return rows


def template_csv(categories: list[str]) -> bytes:
    """CSV cannot hold two sheets, so the notes ride above the table as
    comment lines an admin can delete — the parser skips them anyway."""
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    for row in _instructions_rows(categories):
        w.writerow(["# " + csv_safe(c) for c in row] if row else [])
    w.writerow([])
    w.writerow(COLUMN_NAMES)
    ex = _example_row()
    w.writerow([csv_safe(ex.get(c, "")) for c in COLUMN_NAMES])
    return buf.getvalue().encode("utf-8-sig")


def template_xlsx(categories: list[str]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Items"
    ws.append(COLUMN_NAMES)
    head_fill = PatternFill("solid", fgColor="1F2933")
    for i, name in enumerate(COLUMN_NAMES, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        required = dict((c[0], c[1]) for c in COLUMNS)[name] == "Required"
        ws.column_dimensions[get_column_letter(i)].width = 34 if name in (
            "description", "specifications") else (20 if required else 16)
    ex = _example_row()
    ws.append([csv_safe(ex.get(c, "")) for c in COLUMN_NAMES])
    ws.freeze_panes = "A2"

    notes = wb.create_sheet("Instructions")
    for row in _instructions_rows(categories):
        notes.append([csv_safe(c) for c in row])
    notes.column_dimensions["A"].width = 24
    notes.column_dimensions["B"].width = 12
    notes.column_dimensions["C"].width = 96
    for r in notes.iter_rows():
        for cell in r:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    notes["A1"].font = Font(bold=True, size=13)
    notes["A6"].font = notes["B6"].font = notes["C6"].font = Font(bold=True)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def error_report_csv(rows: list[dict]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow(["row", "id", "name", "field", "problem"])
    for r in rows:
        w.writerow([r.get("row", ""), csv_safe(r.get("id", "")), csv_safe(r.get("name", "")),
                    csv_safe(r.get("field", "")), csv_safe(r.get("message", ""))])
    return buf.getvalue().encode("utf-8-sig")


# ---------------- spreadsheet reading ----------------

def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, bool):
        return "yes" if value else "no"
    return str(value).strip()


def _rows_from_xlsx(data: bytes) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover - depends on the deployment
        raise ImportError_(
            "XLSX support needs the openpyxl package. Install the site's requirements "
            "again, or save the sheet as CSV and upload that.")
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise ImportError_("That file could not be opened as an Excel workbook.")
    try:
        ws = wb["Items"] if "Items" in wb.sheetnames else wb[wb.sheetnames[0]]
        rows = []
        header: list[str] = []
        for excel_row, raw in enumerate(ws.iter_rows(values_only=True), start=1):
            cells = [_cell_text(c) for c in (raw or ())]
            if not header:
                if not any(cells):
                    continue
                header = [_norm_header(c) for c in cells]
                continue
            if not any(cells):
                continue
            rows.append({"_row": excel_row,
                         **{h: c for h, c in zip(header, cells) if h}})
            if len(rows) > MAX_ROWS:
                raise ImportError_(f"That sheet has more than {MAX_ROWS} rows. "
                                   "Please split it into smaller files.")
        if not header:
            raise ImportError_("That sheet has no header row.")
        return rows
    finally:
        wb.close()


def _rows_from_csv(data: bytes) -> list[dict]:
    text = data.decode("utf-8-sig", errors="replace")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows: list[dict] = []
    header: list[str] = []
    for excel_row, cells in enumerate(reader, start=1):
        cells = [c.strip() for c in cells]
        # the template writes its notes as leading "# …" lines
        if not any(cells) or cells[0].startswith("#"):
            continue
        if not header:
            header = [_norm_header(c) for c in cells]
            continue
        rows.append({"_row": excel_row, **{h: c for h, c in zip(header, cells) if h}})
        if len(rows) > MAX_ROWS:
            raise ImportError_(f"That file has more than {MAX_ROWS} rows. "
                               "Please split it into smaller files.")
    if not header:
        raise ImportError_("That file has no header row.")
    return rows


def parse_sheet(data: bytes, filename: str) -> list[dict]:
    if len(data) > MAX_SHEET_BYTES:
        raise ImportError_("The spreadsheet must be 12 MB or smaller.")
    if not data:
        raise ImportError_("That file is empty.")
    name = (filename or "").lower()
    if name.endswith(".xlsx") or data[:2] == b"PK":
        return _rows_from_xlsx(data)
    if name.endswith((".csv", ".txt")):
        return _rows_from_csv(data)
    raise ImportError_("Please upload an .xlsx or .csv file.")


# ---------------- the images ZIP ----------------

class ZipImages:
    """A read-only view of the uploaded archive, keyed by plain file name.

    Nothing is ever extracted to disk: an entry is read into memory only when
    a row asks for it, and only after its declared size has been checked.
    """

    def __init__(self, data: bytes):
        if len(data) > MAX_ZIP_BYTES:
            raise ImportError_("The images ZIP must be 96 MB or smaller.")
        try:
            self._zf = zipfile.ZipFile(io.BytesIO(data))
        except Exception:
            raise ImportError_("That file could not be opened as a ZIP archive.")
        self._by_name: dict[str, str] = {}
        total = 0
        entries = 0
        for info in self._zf.infolist():
            if info.is_dir():
                continue
            entries += 1
            if entries > MAX_ZIP_ENTRIES:
                raise ImportError_(f"The ZIP holds more than {MAX_ZIP_ENTRIES} files.")
            total += info.file_size
            if total > MAX_ZIP_UNCOMPRESSED:
                raise ImportError_("The ZIP expands to more than 320 MB and was rejected.")
            if (info.compress_size > 4096
                    and info.file_size / max(1, info.compress_size) > MAX_ZIP_RATIO):
                raise ImportError_(f"“{self._leaf(info.filename)}” expands far beyond its "
                                   "packed size and was rejected.")
            leaf = self._leaf(info.filename)
            if not leaf or leaf.startswith("."):
                continue          # __MACOSX/, dotfiles, directory-only entries
            if not leaf.lower().endswith((".jpg", ".jpeg", ".png", ".webp")):
                continue          # anything else is simply not offered to a row
            self._by_name.setdefault(leaf.lower(), info.filename)

    @staticmethod
    def _leaf(name: str) -> str:
        """The file's own name, with every directory part and any traversal
        thrown away. A member is only ever addressed by this, so a crafted
        path like ../../etc/passwd cannot name anything outside the archive —
        and is never used to open a file on disk in the first place."""
        flat = str(name or "").replace("\\", "/").split("/")[-1]
        return "" if flat in ("", ".", "..") else flat

    def names(self) -> set[str]:
        return set(self._by_name)

    def read(self, wanted: str) -> bytes:
        key = self._leaf(wanted).lower()
        real = self._by_name.get(key)
        if real is None:
            raise ImportError_(f"“{wanted}” is not in the images ZIP.")
        info = self._zf.getinfo(real)
        if info.file_size > MAX_IMAGE_BYTES:
            raise ImportError_(f"“{wanted}” is larger than 15 MB.")
        with self._zf.open(info) as fh:
            data = fh.read(MAX_IMAGE_BYTES + 1)
        if len(data) > MAX_IMAGE_BYTES:
            raise ImportError_(f"“{wanted}” is larger than 15 MB.")
        return data

    def close(self) -> None:
        try:
            self._zf.close()
        except Exception:
            pass


# ---------------- remote image URLs ----------------

def _public_ip(host: str) -> None:
    """Every address the name resolves to must be a public one. Checking the
    resolved addresses rather than the text of the URL is what stops
    127.0.0.1 wearing a domain name, and refusing when *any* answer is
    private closes the DNS-rebinding gap as far as one process can."""
    try:
        infos = socket.getaddrinfo(host, None)
    except OSError:
        raise ImportError_(f"“{host}” could not be resolved.")
    if not infos:
        raise ImportError_(f"“{host}” could not be resolved.")
    for info in infos:
        ip = ipaddress.ip_address(info[4][0])
        if (ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_multicast
                or ip.is_reserved or ip.is_unspecified):
            raise ImportError_(f"“{host}” points at a private address and was refused.")


def fetch_image_url(url: str) -> bytes:
    import httpx

    parsed = urllib.parse.urlsplit(url)
    if parsed.scheme not in ("http", "https"):
        raise ImportError_("Image links must start with https://")
    if not parsed.hostname:
        raise ImportError_("That image link is not a valid address.")
    if parsed.port not in (None, 80, 443):
        raise ImportError_("Image links may not name a port.")
    _public_ip(parsed.hostname)
    try:
        with httpx.Client(timeout=IMAGE_URL_TIMEOUT_S, follow_redirects=False,
                          trust_env=False) as client:
            with client.stream("GET", url) as res:
                if res.status_code in (301, 302, 303, 307, 308):
                    # a redirect is where an allowed host hands off to a private
                    # one, so it is refused rather than followed
                    raise ImportError_("That image link redirects elsewhere; "
                                       "please use the final address.")
                if res.status_code != 200:
                    raise ImportError_(f"That image link answered {res.status_code}.")
                ctype = (res.headers.get("content-type") or "").split(";")[0].strip().lower()
                if ctype and not ctype.startswith("image/"):
                    raise ImportError_(f"That link is {ctype or 'not an image'}, not an image.")
                buf = bytearray()
                for chunk in res.iter_bytes():
                    buf += chunk
                    if len(buf) > MAX_IMAGE_BYTES:
                        raise ImportError_("That image is larger than 15 MB.")
                return bytes(buf)
    except ImportError_:
        raise
    except Exception:
        raise ImportError_("That image could not be downloaded.")


# ---------------- turning one row into a rental item ----------------

_EXISTING_PATH_RE = re.compile(r"^/(assets|media)/[\w./-]+$")


class _Resolver:
    """Turns whatever an image cell says into a path the rental item can hold,
    once per distinct reference — a picture named by ten rows is fetched and
    stored once."""

    def __init__(self, zip_images: ZipImages | None, by: str):
        self.zip = zip_images
        self.by = by
        self.cache: dict[str, str] = {}
        self.downloads = 0

    def resolve(self, ref: str) -> str:
        ref = str(ref or "").strip()
        if not ref:
            return ""
        if ref in self.cache:
            return self.cache[ref]
        path = self._resolve_uncached(ref)
        self.cache[ref] = path
        return path

    def _resolve_uncached(self, ref: str) -> str:
        from . import media

        if _EXISTING_PATH_RE.match(ref):
            if ref.startswith("/media/"):
                if media.media_file_path(ref[len("/media/"):]) is None:
                    raise ImportError_(f"“{ref}” is not in the Media Library.")
            elif not (config.PUBLIC_DIR / ref.lstrip("/")).is_file():
                raise ImportError_(f"“{ref}” is not a file on the site.")
            return ref
        if "://" in ref or ref.lower().startswith("www."):
            url = ref if "://" in ref else "https://" + ref
            data = fetch_image_url(url)
            self.downloads += 1
            name = urllib.parse.urlsplit(url).path.rsplit("/", 1)[-1] or "image"
        elif self.zip is not None:
            data = self.zip.read(ref)
            name = ref
        else:
            raise ImportError_(
                f"“{ref}” is neither a link nor an existing path, and no images ZIP "
                "was uploaded.")
        try:
            item = media.ingest_library_image(data, name, "", self.by)
        except Exception as exc:
            raise ImportError_(f"“{ref}”: {exc}")
        return "/media/" + item["file"]


# only a real tag — "<" then a letter or "/" — so a description that says
# "fits spaces < 3m" keeps its wording
_TAG_RE = re.compile(r"</?[a-zA-Z][^<>]*>")


def strip_markup(raw: str) -> str:
    """A spreadsheet cell is text, never markup.

    The public pages escape everything they render, so a stored tag is inert
    rather than dangerous — but a name reading "<b>Chair</b>" on the live site
    is still wrong, and a sheet from a supplier is the least trustworthy input
    the panel takes. Tags are removed here, before the row reaches the shared
    rental cleaner, and the preview reports every cell this changed.
    """
    text = str(raw or "")
    if not _TAG_RE.search(text) and "&" not in text:
        return text
    import html as _html

    out = _TAG_RE.sub(" ", text)
    out = _html.unescape(out)
    out = _TAG_RE.sub(" ", out)      # unescaping cannot smuggle a tag back in
    return re.sub(r"\s+", " ", out).strip()


def _split_list(raw: str, sep_extra: str = "") -> list[str]:
    parts = re.split(r"[\n|]+" + (f"|[{sep_extra}]" if sep_extra else ""), str(raw or ""))
    return [p.strip() for p in parts if p.strip()]


def _bool_cell(raw: str, field: str) -> bool:
    v = str(raw or "").strip().lower()
    if v in _TRUE:
        return True
    if v in _FALSE:
        return False
    raise ImportError_(f"{field} must be yes or no (got “{raw}”).")


def _int_cell(raw: str, field: str) -> int:
    v = str(raw or "").strip()
    if not v:
        return 0
    try:
        n = int(float(v))
    except ValueError:
        raise ImportError_(f"{field} must be a whole number (got “{raw}”).")
    if n < 0 or n > 100000:
        raise ImportError_(f"{field} must be between 0 and 100000 (got “{raw}”).")
    return n


def _text_cell(rec: dict, row: dict, col: str) -> str:
    """The cell as text, with any markup taken out and the removal reported."""
    raw = str(row.get(col) or "")
    clean = strip_markup(raw)
    if clean != raw.strip():
        rec["warnings"].append({"field": col, "message": "formatting or HTML removed"})
    return clean


def _has(row: dict, key: str) -> bool:
    return str(row.get(key) or "").strip() != ""


def _cleared(row: dict, key: str) -> bool:
    return str(row.get(key) or "").strip().lower() == CLEAR_TOKEN


# ---------------- validation / preview ----------------

def _existing_categories(products: list[dict]) -> set[str]:
    return {str(p.get("category") or "").strip() for p in products
            if str(p.get("category") or "").strip()}


def validate_rows(rows: list[dict], products: list[dict], resolver: _Resolver,
                  options: dict) -> list[dict]:
    """One entry per spreadsheet row: what it would do, and why not if not.

    Pictures are resolved here rather than at import time on purpose — a
    broken link is something to see in the preview, not a surprise halfway
    through writing five hundred items.
    """
    from . import content

    by_id = {str(p.get("id") or ""): p for p in products}
    known_categories = _existing_categories(products)
    create_categories = bool(options.get("createCategories"))
    on_duplicate = options.get("onDuplicate", "skip")
    image_mode = options.get("imageMode", "append")

    seen_ids: dict[str, int] = {}
    seen_codes: dict[str, int] = {}
    out: list[dict] = []

    for row in rows:
        rec: dict = {"row": row.get("_row"), "errors": [], "warnings": [], "notes": [],
                     "id": str(row.get("id") or "").strip().lower(),
                     "name": str(row.get("name") or "").strip(),
                     "category": str(row.get("category") or "").strip(),
                     "images": 0, "action": "create", "product": None}

        def err(field, message):
            rec["errors"].append({"field": field, "message": message})

        unknown = [k for k in row
                   if k != "_row" and k not in COLUMN_NAMES and str(row.get(k) or "").strip()]
        if unknown:
            rec["warnings"].append({"field": ", ".join(sorted(unknown)[:4]),
                                    "message": "column not recognised — ignored"})

        rid = rec["id"]
        existing = by_id.get(rid)
        if not rid:
            err("id", "id is required.")
        elif not content._SLUG_RE.match(rid):
            err("id", "id must be lowercase letters, digits and dashes, e.g. rent-led-wall.")
        elif rid in seen_ids:
            err("id", f"the same id is on row {seen_ids[rid]} of this sheet.")
        else:
            seen_ids[rid] = rec["row"]

        code = str(row.get("code") or "").strip()
        if code:
            low = code.lower()
            if low in seen_codes:
                rec["warnings"].append(
                    {"field": "code", "message": f"code also used on row {seen_codes[low]}"})
            else:
                seen_codes[low] = rec["row"]
            clash = [p for p in products
                     if str(p.get("code") or "").strip().lower() == low
                     and str(p.get("id") or "") != rid]
            if clash:
                rec["warnings"].append(
                    {"field": "code", "message": f"code already used by “{clash[0].get('name')}”"})

        if existing is not None:
            rec["action"] = "update" if on_duplicate == "update" else "skip"
            rec["duplicate"] = True
            rec["notes"].append("an item with this id already exists")

        # category: never invented quietly
        category = rec["category"]
        if not category and existing is None:
            err("category", "category is required.")
        elif category and category not in known_categories:
            if create_categories:
                rec["warnings"].append({"field": "category",
                                        "message": f"new category “{category}” will be created"})
                known_categories.add(category)
            else:
                err("category", f"category “{category}” does not exist.")

        if not str(row.get("name") or "").strip() and existing is None:
            err("name", "name is required.")

        for field, label in (("stock_ksa", "stock_ksa"), ("stock_uae", "stock_uae")):
            if _has(row, field) and not _cleared(row, field):
                try:
                    _int_cell(row[field], label)
                except ImportError_ as exc:
                    err(field, str(exc))
        if _has(row, "featured") and not _cleared(row, "featured"):
            try:
                _bool_cell(row["featured"], "featured")
            except ImportError_ as exc:
                err("featured", str(exc))

        # images
        refs = [str(row.get(f"image_{i}") or "").strip() for i in range(1, IMAGE_COLUMNS + 1)]
        refs = [r for r in refs if r and r.lower() != CLEAR_TOKEN]
        resolved: list[str] = []
        for i, ref in enumerate(refs, start=1):
            try:
                path = resolver.resolve(ref)
                if path and path not in resolved:
                    resolved.append(path)
            except ImportError_ as exc:
                err(f"image_{i}", str(exc))
            except Exception:
                err(f"image_{i}", "that image could not be read.")
        rec["images"] = len(resolved)

        if rec["errors"]:
            rec["action"] = "error"
            out.append(rec)
            continue
        if rec["action"] == "skip":
            out.append(rec)
            continue

        # build exactly the dict the single-item form posts
        base = dict(existing) if existing else {}
        item: dict = {
            "id": rid,
            "code": base.get("code", ""),
            "name": base.get("name", ""),
            "category": base.get("category", ""),
            "description": base.get("description", ""),
            "tags": list(base.get("tags") or []),
            "specs": list(base.get("specs") or []),
            "featured": bool(base.get("featured")),
            "stockByMarket": dict(base.get("stockByMarket") or {}),
            "images": list(base.get("images") or []),
        }
        text_fields = (("code", "code"), ("name", "name"), ("category", "category"),
                       ("description", "description"))
        for col, key in text_fields:
            if _cleared(row, col):
                item[key] = ""
                rec["notes"].append(f"{col} cleared")
            elif _has(row, col):
                item[key] = _text_cell(rec, row, col)
        if _cleared(row, "specifications"):
            item["specs"] = []
        elif _has(row, "specifications"):
            item["specs"] = _split_list(_text_cell(rec, row, "specifications"))
        if _cleared(row, "search_tags"):
            item["tags"] = []
        elif _has(row, "search_tags"):
            item["tags"] = _split_list(_text_cell(rec, row, "search_tags"), sep_extra=",;")
        if _has(row, "featured") and not _cleared(row, "featured"):
            item["featured"] = _bool_cell(row["featured"], "featured")
        for col, market in (("stock_ksa", "ksa"), ("stock_uae", "uae")):
            if _cleared(row, col):
                item["stockByMarket"][market] = 0
            elif _has(row, col):
                item["stockByMarket"][market] = _int_cell(row[col], col)
            else:
                item["stockByMarket"].setdefault(market, 0)

        if resolved:
            if existing is None or image_mode == "replace":
                item["images"] = resolved
            else:
                item["images"] = list(item["images"]) + [
                    p for p in resolved if p not in item["images"]]
        if any(_cleared(row, f"image_{i}") for i in range(1, IMAGE_COLUMNS + 1)):
            if not resolved:
                item["images"] = []
                rec["notes"].append("images cleared")
        item["images"] = item["images"][:10]
        item["image"] = item["images"][0] if item["images"] else ""

        try:
            clean = content._clean_rental(item)
        except Exception as exc:
            rec["action"] = "error"
            err("", str(exc))
            out.append(rec)
            continue

        # anything the shared cleaner changed is shown rather than done quietly
        for col, key in text_fields:
            if _has(row, col) and clean.get(key) != str(row[col]).strip():
                if re.sub(r"\s+", " ", str(row[col]).strip()) != clean.get(key):
                    rec["warnings"].append({"field": col, "message": "shortened or tidied to fit"})
        rec["product"] = clean
        rec["name"] = clean["name"]
        rec["category"] = clean["category"]
        rec["images"] = len(clean["images"])
        out.append(rec)
    return out


def summarise(results: list[dict]) -> dict:
    return {
        "total": len(results),
        "create": sum(1 for r in results if r["action"] == "create"),
        "update": sum(1 for r in results if r["action"] == "update"),
        "skip": sum(1 for r in results if r["action"] == "skip"),
        "errors": sum(1 for r in results if r["action"] == "error"),
        "warnings": sum(1 for r in results if r["warnings"] and r["action"] != "error"),
        "duplicates": sum(1 for r in results if r.get("duplicate")),
        "images": sum(r["images"] for r in results if r["action"] in ("create", "update")),
    }


# ---------------- staged jobs ----------------

_jobs: dict[str, dict] = {}
_jobs_lock = threading.Lock()
_write_lock = threading.Lock()


def _prune_jobs(now: float) -> None:
    for token, job in list(_jobs.items()):
        if now - job["created"] > STAGE_TTL_S:
            _drop(token)


def _drop(token: str) -> None:
    job = _jobs.pop(token, None)
    if job and job.get("zip") is not None:
        job["zip"].close()
        job["zip"] = None


def stage(rows: list[dict], results: list[dict], filename: str, options: dict,
          zip_images: ZipImages | None, by: str) -> str:
    token = secrets.token_urlsafe(18)
    with _jobs_lock:
        _prune_jobs(time.time())
        _jobs[token] = {
            "created": time.time(), "by": by, "filename": filename,
            "options": options, "results": results, "zip": zip_images,
            "state": "staged", "done": 0,
            "counts": {"created": 0, "updated": 0, "skipped": 0, "failed": 0},
            "failures": [], "started": 0, "finished": 0,
        }
    return token


def get_job(token: str, by: str) -> dict | None:
    with _jobs_lock:
        job = _jobs.get(token)
        if job is None or job["by"] != by:
            return None
        if time.time() - job["created"] > STAGE_TTL_S:
            _drop(token)
            return None
        return job


def job_public(job: dict) -> dict:
    return {
        "state": job["state"], "done": job["done"],
        "total": sum(1 for r in job["results"] if r["action"] in ("create", "update")),
        "counts": dict(job["counts"]), "filename": job["filename"],
        "failures": job["failures"][:200],
    }


def run_job(token: str, by: str) -> None:
    """Write the staged rows, a batch at a time. Each batch is persisted with
    the inventory's own atomic replace, so a failure late in a long import
    leaves the items already written intact rather than rolling the lot back."""
    from . import content

    job = get_job(token, by)
    if job is None:
        return
    with _jobs_lock:
        if job["state"] != "staged":
            return                      # a second click cannot start it twice
        job["state"] = "running"
        job["started"] = int(time.time())

    todo = [r for r in job["results"] if r["action"] in ("create", "update")]
    job["counts"]["skipped"] = sum(1 for r in job["results"] if r["action"] == "skip")
    for r in job["results"]:
        if r["action"] == "error":
            job["counts"]["failed"] += 1
            for e in r["errors"]:
                job["failures"].append({"row": r["row"], "id": r["id"], "name": r["name"],
                                        "field": e["field"], "message": e["message"]})
    try:
        for start in range(0, len(todo), BATCH_SIZE):
            batch = todo[start:start + BATCH_SIZE]
            with _write_lock:
                products, _ = content.rentals_load()
                index = {str(p.get("id") or ""): i for i, p in enumerate(products)}
                applied = 0
                for rec in batch:
                    try:
                        clean = content._clean_rental(rec["product"])
                        if clean["id"] in index:
                            products[index[clean["id"]]] = clean
                            rec["_done"] = "updated"
                        else:
                            index[clean["id"]] = len(products)
                            products.append(clean)
                            rec["_done"] = "created"
                        applied += 1
                    except Exception as exc:
                        rec["_done"] = "failed"
                        job["failures"].append(
                            {"row": rec["row"], "id": rec["id"], "name": rec["name"],
                             "field": "", "message": str(exc)[:300]})
                if applied:
                    content._write_rentals(products)
            for rec in batch:
                done = rec.pop("_done", "failed")
                if done == "created":
                    job["counts"]["created"] += 1
                elif done == "updated":
                    job["counts"]["updated"] += 1
                else:
                    job["counts"]["failed"] += 1
                job["done"] += 1
        job["state"] = "done"
    except Exception as exc:      # pragma: no cover - defensive
        job["state"] = "done"
        job["failures"].append({"row": "", "id": "", "name": "", "field": "",
                                "message": f"import stopped: {exc.__class__.__name__}"})
    finally:
        job["finished"] = int(time.time())
        with _jobs_lock:
            if job.get("zip") is not None:
                job["zip"].close()
                job["zip"] = None
        record_history(job)


# ---------------- history ----------------

def _ensure() -> None:
    from . import adminauth as aa

    with aa._lock:
        conn = aa._connect()
        conn.executescript("""
            CREATE TABLE IF NOT EXISTS rental_imports (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                ts INTEGER NOT NULL,
                user_email TEXT NOT NULL DEFAULT '',
                filename TEXT NOT NULL DEFAULT '',
                total INTEGER NOT NULL DEFAULT 0,
                created INTEGER NOT NULL DEFAULT 0,
                updated INTEGER NOT NULL DEFAULT 0,
                skipped INTEGER NOT NULL DEFAULT 0,
                failed INTEGER NOT NULL DEFAULT 0,
                failures TEXT NOT NULL DEFAULT '[]'
            );
        """)
        conn.commit()


def record_history(job: dict) -> None:
    from . import adminauth as aa

    _ensure()
    counts = job["counts"]
    # the report is kept, the uploaded file is not
    failures = json.dumps(job["failures"][:500], ensure_ascii=False)
    with aa._lock:
        conn = aa._connect()
        conn.execute(
            "INSERT INTO rental_imports (ts, user_email, filename, total, created, updated, "
            "skipped, failed, failures) VALUES (?,?,?,?,?,?,?,?,?)",
            (int(time.time()), job["by"][:200], job["filename"][:200], len(job["results"]),
             counts["created"], counts["updated"], counts["skipped"], counts["failed"],
             failures))
        conn.execute(
            "DELETE FROM rental_imports WHERE id NOT IN "
            "(SELECT id FROM rental_imports ORDER BY id DESC LIMIT ?)", (KEEP_HISTORY,))
        conn.commit()


def history(limit: int = 20) -> list[dict]:
    from . import adminauth as aa

    _ensure()
    rows = aa._connect().execute(
        "SELECT id, ts, user_email, filename, total, created, updated, skipped, failed "
        "FROM rental_imports ORDER BY id DESC LIMIT ?", (max(1, min(100, limit)),)).fetchall()
    return [dict(r) for r in rows]


def history_failures(entry_id: int) -> list[dict] | None:
    from . import adminauth as aa

    _ensure()
    row = aa._connect().execute(
        "SELECT failures FROM rental_imports WHERE id=?", (entry_id,)).fetchone()
    if row is None:
        return None
    try:
        return json.loads(row["failures"])
    except ValueError:
        return []
